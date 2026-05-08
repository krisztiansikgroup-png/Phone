"""
Generate 'Foaie Colectivă de Prezență' Excel workbook.

Format matches SC Piese Transilvane SRL template exactly:
  Row 1 : Company name
  Row 2 : Document title
  Row 3 : Month / working-days count / TOTAL ORE header
  Row 4 : Nr. | Nume și Prenume | ORE ZILNICE (spanning days)
  Row 5 : Day numbers 1-31
  Rows 6+: Employee rows  (8 | X | Co | Bo | D | S | NN)
  Last 5 : Legend
"""
import calendar
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from holidays_ro import get_public_holidays

# ── status mapping (our codes → foaie codes) ─────────────────────────────────
STATUS_TO_FOAIE = {
    "P":   "8",
    "LS":  "X",
    "SL":  "X",
    "CO":  "Co",
    "CM":  "Bo",
    "CFP": "NN",
    "CIC": "S",
    "DO":  "D",
    "REC": "8",
    "NN":  "NN",
}

_MONTH_NAMES = [
    "Ianuarie", "Februarie", "Martie", "Aprilie", "Mai", "Iunie",
    "Iulie", "August", "Septembrie", "Octombrie", "Noiembrie", "Decembrie",
]

_THIN = Side(style="thin")
_MEDIUM = Side(style="medium")
_ALL = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_THICK = Border(left=_MEDIUM, right=_MEDIUM, top=_MEDIUM, bottom=_MEDIUM)

_FILL_HEADER = PatternFill("solid", fgColor="BDD7EE")   # blue header
_FILL_WEEKEND = PatternFill("solid", fgColor="D9D9D9")  # grey for X
_FILL_CO = PatternFill("solid", fgColor="FFFF00")       # yellow
_FILL_BO = PatternFill("solid", fgColor="FF9900")       # orange
_FILL_D = PatternFill("solid", fgColor="92D050")        # green
_FILL_S = PatternFill("solid", fgColor="FF0000")        # red
_FILL_NN = PatternFill("solid", fgColor="FF0000")       # red
_FILL_TOTAL = PatternFill("solid", fgColor="E2EFDA")    # light green

_FONT_TITLE = Font(name="Arial", size=12, bold=True)
_FONT_HEADER = Font(name="Arial", size=9, bold=True)
_FONT_BODY = Font(name="Arial", size=9)
_FONT_LEGEND = Font(name="Arial", size=8, italic=True)

_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center")


def _c(ws, row, col, value="", font=None, fill=None, border=True, align=None):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font or _FONT_BODY
    if fill:
        cell.fill = fill
    if border:
        cell.border = _ALL
    cell.alignment = align or _CENTER
    return cell


def generate_foaie(year: int, month: int, employees: list,
                   records_by_emp: dict) -> bytes:
    """
    employees       : list of Employee objects (ordered)
    records_by_emp  : {employee_id: {date: status_code}}
    Returns raw .xlsx bytes in Foaie Colectivă de Prezență format.
    """
    wb = Workbook()
    ws = wb.active
    month_name = _MONTH_NAMES[month - 1]
    ws.title = f"{year}-{month_name[:3]}"

    _, num_days = calendar.monthrange(year, month)
    holidays = get_public_holidays(year)
    all_days = [date(year, month, d) for d in range(1, num_days + 1)]

    # count working days
    working_days = sum(
        1 for d in all_days
        if d.weekday() < 5 and d not in holidays
    )

    # Column layout: A=Nr, B=Nume, C=blank, D..D+num_days-1=days, last=Total
    day_col_start = 4       # column D
    total_col = day_col_start + num_days   # after all day columns
    last_col = total_col

    # ── Row 1: Company ────────────────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c = ws.cell(row=1, column=1,
                value="Unitatea: S.C. PIESE TRANSILVANE S.R.L.")
    c.font = _FONT_TITLE
    c.alignment = _LEFT
    ws.row_dimensions[1].height = 18

    # ── Row 2: Document title ─────────────────────────────────────────────────
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c = ws.cell(row=2, column=1, value="FOAIE COLECTIVĂ DE PREZENȚĂ")
    c.font = _FONT_TITLE
    c.alignment = _CENTER
    ws.row_dimensions[2].height = 18

    # ── Row 3: Month / working days / TOTAL ORE ───────────────────────────────
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=3)
    c = ws.cell(row=3, column=1,
                value=f"Pe luna: {month_name.upper()} {year}")
    c.font = _FONT_HEADER
    c.alignment = _LEFT

    # "N zile lucrătoare" near the middle
    mid_col = day_col_start + num_days // 2
    ws.merge_cells(start_row=3, start_column=day_col_start,
                   end_row=3, end_column=mid_col)
    c = ws.cell(row=3, column=day_col_start,
                value=f"{working_days} zile lucrătoare")
    c.font = _FONT_HEADER
    c.alignment = _CENTER

    ws.cell(row=3, column=last_col,
            value="TOTAL ORE").font = _FONT_HEADER
    ws.cell(row=3, column=last_col).alignment = _CENTER
    ws.row_dimensions[3].height = 16

    # ── Row 4: Column headers ─────────────────────────────────────────────────
    _c(ws, 4, 1, "Nr.", font=_FONT_HEADER, fill=_FILL_HEADER)
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=3)
    _c(ws, 4, 2, "Nume și Prenume", font=_FONT_HEADER, fill=_FILL_HEADER,
       align=_CENTER)
    ws.merge_cells(start_row=4, start_column=day_col_start,
                   end_row=4, end_column=total_col - 1)
    _c(ws, 4, day_col_start, "ORE ZILNICE",
       font=_FONT_HEADER, fill=_FILL_HEADER, align=_CENTER)
    _c(ws, 4, total_col, "", font=_FONT_HEADER, fill=_FILL_HEADER)
    ws.row_dimensions[4].height = 16

    # ── Row 5: Day numbers ────────────────────────────────────────────────────
    _c(ws, 5, 1, "", font=_FONT_HEADER)
    _c(ws, 5, 2, "", font=_FONT_HEADER)
    _c(ws, 5, 3, "", font=_FONT_HEADER)
    for di, d in enumerate(all_days):
        col = day_col_start + di
        is_off = d.weekday() >= 5 or d in holidays
        fill = _FILL_WEEKEND if is_off else _FILL_HEADER
        _c(ws, 5, col, d.day, font=_FONT_HEADER, fill=fill)
    _c(ws, 5, total_col, "", font=_FONT_HEADER, fill=_FILL_HEADER)
    ws.row_dimensions[5].height = 14

    # ── Employee rows ─────────────────────────────────────────────────────────
    for row_idx, emp in enumerate(employees):
        row = 6 + row_idx
        emp_records = records_by_emp.get(emp.id, {})

        _c(ws, row, 1, f"{row_idx + 1}.", font=_FONT_BODY)
        ws.merge_cells(start_row=row, start_column=2,
                       end_row=row, end_column=3)
        _c(ws, row, 2, emp.name, font=_FONT_BODY, align=_LEFT)

        total_hours = 0
        for di, d in enumerate(all_days):
            col = day_col_start + di
            is_weekend = d.weekday() >= 5
            is_holiday = d in holidays

            if is_weekend:
                default = "LS"
            elif is_holiday:
                default = "SL"
            else:
                default = "P"

            status = emp_records.get(d, default)
            foaie_val = STATUS_TO_FOAIE.get(status, "8")

            if foaie_val == "8":
                total_hours += 8
                fill = None
            elif foaie_val == "X":
                fill = _FILL_WEEKEND
            elif foaie_val == "Co":
                fill = _FILL_CO
            elif foaie_val == "Bo":
                fill = _FILL_BO
            elif foaie_val == "D":
                fill = _FILL_D
            elif foaie_val in ("S", "NN"):
                fill = _FILL_S
            else:
                fill = None

            display = foaie_val if foaie_val != "8" else 8
            _c(ws, row, col, display, font=_FONT_BODY, fill=fill)

        _c(ws, row, total_col, total_hours,
           font=Font(name="Arial", size=9, bold=True), fill=_FILL_TOTAL)
        ws.row_dimensions[row].height = 15

    # ── Legend ────────────────────────────────────────────────────────────────
    legend_start = 6 + len(employees)
    legend = [
        "Co-concediu de odihnă",
        "Bo-boală",
        "D-delegație",
        "S-suspendat",
        "NN-nelucrat, neplătit",
    ]
    for li, text in enumerate(legend):
        row = legend_start + li
        ws.merge_cells(start_row=row, start_column=2,
                       end_row=row, end_column=last_col)
        c = ws.cell(row=row, column=2, value=text)
        c.font = _FONT_LEGEND
        c.alignment = _LEFT
        ws.row_dimensions[row].height = 13

    # ── Column widths ─────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 2
    for di in range(num_days):
        ws.column_dimensions[get_column_letter(day_col_start + di)].width = 3.5
    ws.column_dimensions[get_column_letter(total_col)].width = 9

    ws.freeze_panes = "D6"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_annual_foaie(year: int, employees: list,
                           records_by_month: dict) -> bytes:
    """Generate a full-year workbook with one sheet per month."""
    wb = Workbook()
    wb.remove(wb.active)   # remove default empty sheet

    for month in range(1, 13):
        month_name = _MONTH_NAMES[month - 1]
        ws = wb.create_sheet(title=f"{year}-{month_name[:3]}")
        _fill_sheet(ws, year, month, employees,
                    records_by_month.get(month, {}))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fill_sheet(ws, year, month, employees, records_by_emp):
    """Internal: write one month's data into an existing worksheet."""
    from openpyxl import Workbook as _WB
    _tmp = _WB()
    _tmp_ws = _tmp.active
    data_bytes = generate_foaie(year, month, employees, records_by_emp)
    from openpyxl import load_workbook
    from io import BytesIO as _BIO
    tmp_wb = load_workbook(_BIO(data_bytes))
    tmp_ws = tmp_wb.active
    # copy cell by cell (simple approach)
    for row in tmp_ws.iter_rows():
        for cell in row:
            target = ws.cell(row=cell.row, column=cell.column,
                             value=cell.value)
            if cell.font:
                target.font = cell.font.copy()
            if cell.fill and cell.fill.fill_type != "none":
                target.fill = cell.fill.copy()
            if cell.border:
                target.border = cell.border.copy()
            if cell.alignment:
                target.alignment = cell.alignment.copy()
    for col_letter, dim in tmp_ws.column_dimensions.items():
        ws.column_dimensions[col_letter].width = dim.width
    for row_num, dim in tmp_ws.row_dimensions.items():
        ws.row_dimensions[row_num].height = dim.height
